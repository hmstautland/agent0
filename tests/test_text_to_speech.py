import re
from pathlib import Path

import numpy as np
import pytest

import core.text_to_speech as tts


# --- validation --------------------------------------------------------

def test_validate_speed_accepts_boundaries_and_numeric_strings():
    assert tts.validate_speed(1.0) == 1.0
    assert tts.validate_speed("0.5") == 0.5
    assert tts.validate_speed(2.0) == 2.0


@pytest.mark.parametrize("speed", [0.4, 2.1, "fast", None])
def test_validate_speed_rejects_out_of_range_and_non_numeric(speed):
    with pytest.raises(ValueError):
        tts.validate_speed(speed)


@pytest.mark.parametrize("fmt", ["wav", "mp3", "flac", "ogg", "MP3"])
def test_validate_output_format_accepts_known_formats_case_insensitively(fmt):
    assert tts.validate_output_format(fmt) == fmt.lower()


def test_validate_output_format_defaults_when_omitted():
    assert tts.validate_output_format(None) == tts.DEFAULT_OUTPUT_FORMAT


def test_validate_output_format_rejects_unknown():
    with pytest.raises(ValueError):
        tts.validate_output_format("aac")


# --- voice list ----------------------------------------------------------

def test_get_voices_only_has_british_english_group():
    groups = {voice["group"] for voice in tts.get_voices()}
    assert groups == {"British English"}


def test_get_voices_includes_default_voice():
    assert any(v["id"] == tts.DEFAULT_VOICE for v in tts.get_voices())


# --- chunking --------------------------------------------------------------

def test_chunk_text_keeps_short_text_as_one_chunk():
    assert tts._chunk_text("Hello there.", max_chars=2000) == ["Hello there."]


def test_chunk_text_splits_long_text_and_preserves_content():
    text = "One sentence. " * 500
    chunks = tts._chunk_text(text, max_chars=80)

    assert len(chunks) > 1
    assert all(len(chunk) <= 80 for chunk in chunks)
    assert " ".join(chunks).replace("  ", " ").strip() == text.strip()


def test_chunk_text_handles_text_with_no_sentence_punctuation():
    text = "just one long run-on clause with no punctuation at all"
    assert tts._chunk_text(text, max_chars=2000) == [text]


# --- input validation on the public entry point -----------------------------

def test_synthesize_rejects_empty_text():
    with pytest.raises(ValueError, match="No text"):
        tts.synthesize("")


def test_synthesize_rejects_text_over_word_limit():
    text = "word " * (tts.MAX_TTS_WORDS + 1)
    with pytest.raises(ValueError, match="too long"):
        tts.synthesize(text)


def test_synthesize_rejects_unknown_voice():
    with pytest.raises(ValueError, match="Unknown voice"):
        tts.synthesize("hello", voice_id="xx_not_a_real_voice")


def test_synthesize_rejects_invalid_speed_before_calling_kokoro(monkeypatch):
    def fail_if_called(*args, **kwargs):
        raise AssertionError("Kokoro should not be reached when speed is invalid")

    monkeypatch.setattr(tts, "_synthesize_kokoro", fail_if_called)

    with pytest.raises(ValueError):
        tts.synthesize("hello", voice_id="bf_isabella", speed=10)


# --- Kokoro synthesis, fully mocked (no model weights involved) ------------

def test_synthesize_kokoro_reuses_a_cached_pipeline_per_language(monkeypatch):
    constructed = []

    class FakePipeline:
        def __call__(self, text, voice=None, speed=1):
            yield "gs", "ps", np.zeros(10, dtype="float32")

    def fake_get_pipeline(lang_code):
        if lang_code not in tts._pipelines:
            constructed.append(lang_code)
            tts._pipelines[lang_code] = FakePipeline()
        return tts._pipelines[lang_code]

    monkeypatch.setattr(tts, "_pipelines", {})
    monkeypatch.setattr(tts, "_resolve_device", lambda: "cpu")
    monkeypatch.setattr(tts, "_get_pipeline", fake_get_pipeline)

    tts._synthesize_kokoro("Hi there.", "bf_isabella", 1.0)
    tts._synthesize_kokoro("Hi again.", "bf_isabella", 1.0)

    assert constructed == ["b"]  # constructed once, reused on the second call


def test_synthesize_kokoro_concatenates_chunk_audio(monkeypatch):
    monkeypatch.setattr(tts, "_chunk_text", lambda text, max_chars=None: ["chunk one", "chunk two"])

    def fake_pipeline(lang_code):
        def call(text, voice=None, speed=1):
            yield None, None, np.full(5, len(text), dtype="float32")

        return call

    monkeypatch.setattr(tts, "_get_pipeline", fake_pipeline)

    result = tts._synthesize_kokoro("irrelevant", "bf_isabella", 1.0)

    assert result.shape == (10,)


def test_synthesize_full_flow_without_kokoro_installed(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)
    monkeypatch.setattr(
        tts, "_synthesize_kokoro", lambda text, voice_id, speed: np.zeros(1200, dtype="float32")
    )

    result = tts.synthesize("Testing without real kokoro.", voice_id="bf_emma", speed=1.25, output_format="wav")

    assert Path(result["path"]).exists()
    assert result["mime_type"] == "audio/wav"


def test_synthesize_reports_missing_kokoro_dependency_clearly(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)

    def raise_import_error(text, voice_id, speed):
        raise ImportError("no module named kokoro")

    monkeypatch.setattr(tts, "_synthesize_kokoro", raise_import_error)

    with pytest.raises(RuntimeError, match="Kokoro TTS is not installed"):
        tts.synthesize("hello", voice_id="bf_isabella")


# --- WAV/FLAC output, using the real soundfile encoder ----------------------

def test_synthesize_generates_a_valid_wav_file(monkeypatch, tmp_path):
    import wave

    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)
    samples = np.linspace(-0.5, 0.5, 2400, dtype="float32")
    monkeypatch.setattr(tts, "_synthesize_kokoro", lambda text, voice_id, speed: samples)

    result = tts.synthesize("Hello world.", voice_id="bf_isabella", output_format="wav")

    out_path = Path(result["path"])
    assert out_path.exists()

    with wave.open(str(out_path), "rb") as wav_file:
        assert wav_file.getframerate() == tts.SAMPLE_RATE
        assert wav_file.getnframes() == len(samples)


# --- user-chosen filename / title ------------------------------------------

def test_build_filename_stem_sanitizes_a_given_title(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)

    assert tts._build_filename_stem("My Weird/Title: v2?") == "My Weird_Title_ v2"


def test_build_filename_stem_defaults_to_a_date_time_pattern_when_blank(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)

    stem = tts._build_filename_stem("   ")

    assert re.match(r"^tts_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}$", stem)


def test_build_filename_stem_dedupes_against_an_existing_file(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)
    (tmp_path / "My Title.wav").touch()

    assert tts._build_filename_stem("My Title") == "My Title_2"


def test_synthesize_uses_the_given_title_as_the_filename(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)
    monkeypatch.setattr(tts, "_synthesize_kokoro", lambda text, voice_id, speed: np.zeros(100, dtype="float32"))

    result = tts.synthesize("hello", voice_id="bf_isabella", output_format="wav", title="Weekly Update")

    assert result["filename"] == "Weekly Update.wav"


def test_synthesize_falls_back_to_default_name_without_a_title(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)
    monkeypatch.setattr(tts, "_synthesize_kokoro", lambda text, voice_id, speed: np.zeros(100, dtype="float32"))

    result = tts.synthesize("hello", voice_id="bf_isabella", output_format="wav")

    assert re.match(r"^tts_\d{4}-\d{2}-\d{2}_\d{2}-\d{2}\.wav$", result["filename"])


def test_encode_flac_round_trips_the_same_audio(tmp_path):
    import soundfile as sf

    wav_path = tmp_path / "sample.wav"
    samples = np.linspace(-0.5, 0.5, 4800, dtype="float32")
    sf.write(str(wav_path), samples, tts.SAMPLE_RATE, subtype="PCM_16")

    flac_path = tts._encode_flac(wav_path)

    assert flac_path.suffix == ".flac"
    assert flac_path.exists()

    data, sr = sf.read(str(flac_path))
    assert sr == tts.SAMPLE_RATE
    assert len(data) == len(samples)


# --- two-speaker dialogue ----------------------------------------------------

def test_parse_dialogue_splits_labeled_lines():
    text = "A: Hello there.\nB: Hi, how are you?\nA: Doing well, thanks."
    assert tts.parse_dialogue(text) == [
        ("A", "Hello there."),
        ("B", "Hi, how are you?"),
        ("A", "Doing well, thanks."),
    ]


def test_parse_dialogue_is_case_insensitive_and_skips_blank_lines():
    text = "a: hello\n\nb: hi\n"
    assert tts.parse_dialogue(text) == [("A", "hello"), ("B", "hi")]


def test_parse_dialogue_merges_consecutive_same_speaker_lines():
    text = "A: First line.\nA: Second line.\nB: Reply."
    assert tts.parse_dialogue(text) == [("A", "First line. Second line."), ("B", "Reply.")]


def test_parse_dialogue_rejects_lines_without_a_speaker_label():
    with pytest.raises(ValueError, match='"A:" or "B:"'):
        tts.parse_dialogue("A: Hello.\nJust a plain line.")


def test_parse_dialogue_rejects_empty_input():
    with pytest.raises(ValueError, match="No dialogue lines"):
        tts.parse_dialogue("   \n  ")


def test_synthesize_dialogue_rejects_unknown_voice():
    with pytest.raises(ValueError, match="Unknown voice"):
        tts.synthesize_dialogue("A: hi\nB: hi", voice_a="xx_bogus", voice_b="bm_george")


def test_synthesize_dialogue_calls_kokoro_once_per_turn_with_the_right_voice(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)
    calls = []

    def fake_synthesize_kokoro(text, voice_id, speed):
        calls.append((text, voice_id))
        return np.zeros(100, dtype="float32")

    monkeypatch.setattr(tts, "_synthesize_kokoro", fake_synthesize_kokoro)

    result = tts.synthesize_dialogue(
        "A: Hello there.\nB: Hi, how are you?",
        voice_a="bf_isabella",
        voice_b="bm_george",
        output_format="wav",
    )

    assert calls == [
        ("Hello there.", "bf_isabella"),
        ("Hi, how are you?", "bm_george"),
    ]
    assert Path(result["path"]).exists()


def test_synthesize_dialogue_inserts_a_silence_gap_between_turns(monkeypatch, tmp_path):
    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)
    monkeypatch.setattr(tts, "_synthesize_kokoro", lambda text, voice_id, speed: np.ones(100, dtype="float32"))

    result = tts.synthesize_dialogue("A: one\nB: two\nA: three", output_format="wav")

    import wave

    with wave.open(result["path"], "rb") as wav_file:
        expected_gap_frames = int(tts.DIALOGUE_TURN_GAP_SECONDS * tts.SAMPLE_RATE)
        expected_total = 3 * 100 + 2 * expected_gap_frames
        assert wav_file.getnframes() == expected_total


# --- standalone file converter (drag-and-drop, no synthesis involved) -------

def test_convert_audio_file_rejects_unsupported_extension(monkeypatch, tmp_path):
    import io

    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)

    with pytest.raises(ValueError, match="Unsupported input file type"):
        tts.convert_audio_file(io.BytesIO(b"not audio"), "malware.exe", "wav")


def test_convert_audio_file_copies_without_ffmpeg_when_format_unchanged(monkeypatch, tmp_path):
    import io

    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)

    def fail_if_called(*args, **kwargs):
        raise AssertionError("ffmpeg should not run when input and output formats match")

    monkeypatch.setattr(tts.subprocess, "run", fail_if_called)

    content = b"RIFF....WAVEfmt fake wav bytes"
    result = tts.convert_audio_file(io.BytesIO(content), "clip.wav", "wav")

    assert Path(result["path"]).read_bytes() == content
    assert result["mime_type"] == "audio/wav"


def test_convert_audio_file_rejects_uploads_over_the_size_cap(monkeypatch, tmp_path):
    import io

    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)
    monkeypatch.setattr(tts, "MAX_CONVERT_FILE_BYTES", 10)

    with pytest.raises(ValueError, match="too large"):
        tts.convert_audio_file(io.BytesIO(b"way more than ten bytes of content"), "clip.wav", "wav")

    assert list(tmp_path.iterdir()) == []  # the oversized partial upload is cleaned up


def test_convert_audio_file_real_ffmpeg_conversion(tmp_path, monkeypatch):
    import soundfile as sf

    monkeypatch.setattr(tts, "AUDIO_DIR", tmp_path)

    source_wav = tmp_path / "source.wav"
    samples = np.linspace(-0.5, 0.5, 4800, dtype="float32")
    sf.write(str(source_wav), samples, tts.SAMPLE_RATE, subtype="PCM_16")

    with open(source_wav, "rb") as f:
        result = tts.convert_audio_file(f, "source.wav", "flac")

    assert result["path"].endswith(".flac")
    data, sr = sf.read(result["path"])
    assert sr == tts.SAMPLE_RATE
    assert len(data) == len(samples)


# --- document text extraction (drag-and-drop onto the text box) ------------

def test_extract_text_from_document_reads_plain_txt():
    import io

    text = tts.extract_text_from_document(io.BytesIO(b"Hello from a text file."), "notes.txt")

    assert text == "Hello from a text file."


def test_extract_text_from_document_reads_docx(tmp_path):
    import docx

    docx_path = tmp_path / "notes.docx"
    document = docx.Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    document.save(str(docx_path))

    with open(docx_path, "rb") as f:
        text = tts.extract_text_from_document(f, "notes.docx")

    assert text == "First paragraph.\nSecond paragraph."


def test_extract_text_from_document_rejects_legacy_doc():
    import io

    with pytest.raises(ValueError, match=r"\.doc files aren't supported"):
        tts.extract_text_from_document(io.BytesIO(b"whatever"), "report.doc")


def test_extract_text_from_document_rejects_unsupported_extension():
    import io

    with pytest.raises(ValueError, match="Unsupported document type"):
        tts.extract_text_from_document(io.BytesIO(b"whatever"), "report.pdf")


def test_extract_text_from_document_rejects_empty_file():
    import io

    with pytest.raises(ValueError, match="No text found"):
        tts.extract_text_from_document(io.BytesIO(b"   \n  "), "empty.txt")


def test_extract_text_from_document_rejects_corrupt_docx():
    import io

    with pytest.raises(ValueError, match="Could not read that .docx file"):
        tts.extract_text_from_document(io.BytesIO(b"not a real docx"), "fake.docx")


# --- markdown stripping (.md files only - .txt/pasted text is untouched) ---

def test_strip_markdown_for_speech_strips_headings():
    assert tts._strip_markdown_for_speech("# Title\n## Subtitle\ntext") == "Title\nSubtitle\ntext"


def test_strip_markdown_for_speech_does_not_touch_a_hashtag():
    # No space after '#' - not a valid ATX heading, so it's left alone.
    assert tts._strip_markdown_for_speech("#TeamWork is great") == "#TeamWork is great"


def test_strip_markdown_for_speech_strips_bold_and_italic():
    assert tts._strip_markdown_for_speech("**bold** and *italic* and _also italic_") == "bold and italic and also italic"


def test_strip_markdown_for_speech_strips_inline_code_and_links():
    assert tts._strip_markdown_for_speech("See `core/llm.py` or [the docs](https://example.com).") == (
        "See core/llm.py or the docs."
    )


def test_strip_markdown_for_speech_strips_code_fences():
    text = "before\n```python\nprint('hi')\n```\nafter"
    assert tts._strip_markdown_for_speech(text) == "before\nafter"


def test_strip_markdown_for_speech_strips_list_markers_and_blockquotes():
    text = "- first\n- second\n> quoted"
    assert tts._strip_markdown_for_speech(text) == "first\nsecond\nquoted"


def test_strip_markdown_for_speech_strips_horizontal_rules():
    assert tts._strip_markdown_for_speech("above\n---\nbelow") == "above\nbelow"


def test_strip_markdown_for_speech_leaves_email_addresses_alone():
    assert tts._strip_markdown_for_speech("Contact me at user@example.com please.") == (
        "Contact me at user@example.com please."
    )


def test_extract_text_from_document_strips_markdown_for_md_files():
    import io

    text = tts.extract_text_from_document(io.BytesIO(b"# Architecture\n\nSee `agent.py` for details."), "notes.md")

    assert text == "Architecture\n\nSee agent.py for details."
